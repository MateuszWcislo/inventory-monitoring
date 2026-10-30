import csv
import datetime
from decimal import Decimal, ROUND_HALF_UP
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from django.shortcuts import render
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse, Http404
from django.db.models import Sum, Count, F, Q, DecimalField, Value
from django.db.models.functions import Coalesce
from django.utils import timezone

from inventory.models import Product, ProductBatch, Location
from orders.models import Order
from work_orders.models import WorkOrder, WorkOrderProduct
from suppliers.models import Supplier


def _parse_date_range(request):
    """Pomocnicze parsowanie zakresu dat (data od - data do) oraz presetów czasowych."""
    preset = request.GET.get('preset', '').strip()
    date_from_str = request.GET.get('date_from', '').strip()
    date_to_str = request.GET.get('date_to', '').strip()

    today = timezone.now().date()
    date_from = None
    date_to = None

    if preset == '7d':
        date_from = today - datetime.timedelta(days=7)
        date_to = today
    elif preset == '30d':
        date_from = today - datetime.timedelta(days=30)
        date_to = today
    elif preset == '90d':
        date_from = today - datetime.timedelta(days=90)
        date_to = today
    elif preset == 'month':
        date_from = today.replace(day=1)
        date_to = today
    elif preset == 'all':
        date_from = None
        date_to = None
    else:
        if date_from_str:
            try:
                date_from = datetime.datetime.strptime(date_from_str, '%Y-%m-%d').date()
            except ValueError:
                date_from = None
        if date_to_str:
            try:
                date_to = datetime.datetime.strptime(date_to_str, '%Y-%m-%d').date()
            except ValueError:
                date_to = None

    return date_from, date_to, preset


def _get_estimated_product_cost(product, tenant):
    """Pomocnicza funkcja do wyznaczenia szacunkowej ceny zakupu produktu."""
    latest_batch = product.batches.filter(tenant=tenant).order_by('-created_at').first()
    if latest_batch and latest_batch.net_price > 0:
        return latest_batch.net_price, latest_batch.gross_price

    return Decimal('0.00'), Decimal('0.00')


def _calculate_restock_report_data(tenant, location_id=None):
    """Generuje dane raportu zapotrzebowania i przewidywanych kosztów zamówień."""
    products_qs = Product.objects.filter(tenant=tenant).select_related('tenant').prefetch_related(
        'batches', 'supplier_mappings__supplier'
    )

    report_items = []
    total_shortage_qty = 0
    total_estimated_cost_net = Decimal('0.00')
    total_estimated_cost_gross = Decimal('0.00')

    for prod in products_qs:
        batches_qs = prod.batches.filter(tenant=tenant)
        if location_id:
            batches_qs = batches_qs.filter(location_id=location_id)

        current_stock = sum(b.current_stock or 0 for b in batches_qs)
        min_threshold = prod.min_threshold or 0

        # Weryfikacja czy stan jest poniżej minimum
        if current_stock < min_threshold:
            shortage = min_threshold - current_stock
            net_unit_price, gross_unit_price = _get_estimated_product_cost(prod, tenant)
            
            item_total_net = (net_unit_price * Decimal(shortage)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
            item_total_gross = (gross_unit_price * Decimal(shortage)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

            primary_supplier = prod.supplier_mappings.filter(tenant=tenant).first()
            supplier_name = primary_supplier.supplier.name if primary_supplier else "Brak powiązania"

            report_items.append({
                'product': prod,
                'current_stock': current_stock,
                'min_stock_level': min_threshold,
                'shortage_quantity': shortage,
                'unit_price_net': net_unit_price,
                'unit_price_gross': gross_unit_price,
                'total_cost_net': item_total_net,
                'total_cost_gross': item_total_gross,
                'supplier_name': supplier_name,
            })

            total_shortage_qty += shortage
            total_estimated_cost_net += item_total_net
            total_estimated_cost_gross += item_total_gross

    report_items.sort(key=lambda x: x['total_cost_net'], reverse=True)

    return {
        'items': report_items,
        'total_products_count': len(report_items),
        'total_shortage_qty': total_shortage_qty,
        'total_estimated_cost_net': total_estimated_cost_net,
        'total_estimated_cost_gross': total_estimated_cost_gross,
    }


@login_required
def analytics_dashboard(request):
    """Główny dashboard analityczny z podsumowaniami, oknami czasowymi i wykresami."""
    tenant = request.user.tenant

    # Parsowanie zakresu dat
    date_from, date_to, preset = _parse_date_range(request)

    # 1. Podstawowe bieżące metryki magazynowe (stan aktualny)
    total_products = Product.objects.filter(tenant=tenant).count()
    
    batches = ProductBatch.objects.filter(tenant=tenant)
    total_stock_items = batches.aggregate(total=Coalesce(Sum('current_stock'), 0))['total']
    
    batches_with_val = batches.annotate(
        batch_val_net=F('current_stock') * F('net_price'),
        batch_val_gross=F('current_stock') * F('gross_price')
    )
    total_inventory_net = batches_with_val.aggregate(val=Coalesce(Sum('batch_val_net'), Decimal('0.00'), output_field=DecimalField()))['val']
    total_inventory_gross = batches_with_val.aggregate(val=Coalesce(Sum('batch_val_gross'), Decimal('0.00'), output_field=DecimalField()))['val']

    # 2. Braki magazynowe
    restock_data = _calculate_restock_report_data(tenant)
    low_stock_products_count = restock_data['total_products_count']
    estimated_restock_cost = restock_data['total_estimated_cost_net']

    # 3. Zastosowanie filtrów dat do zamówień (Orders)
    orders_qs = Order.objects.filter(tenant=tenant)
    if date_from:
        orders_qs = orders_qs.filter(created_at__date__gte=date_from)
    if date_to:
        orders_qs = orders_qs.filter(created_at__date__lte=date_to)

    # Metryki zamówień w wybranym oknie czasowym
    period_orders_count = orders_qs.count()
    completed_or_ordered = orders_qs.filter(status__in=['ORDERED', 'COMPLETED'])
    period_ordered_items_qty = completed_or_ordered.aggregate(total=Coalesce(Sum('quantity'), 0))['total']
    period_ordered_val_net = completed_or_ordered.annotate(
        val=F('quantity') * F('net_price')
    ).aggregate(total=Coalesce(Sum('val'), Decimal('0.00'), output_field=DecimalField()))['total']

    # Wykres: Top najczęściej zamawianych produktów w wybranym oknie
    top_ordered_qs = (
        completed_or_ordered
        .values('product__name', 'product_name_snapshot')
        .annotate(
            total_ordered_qty=Coalesce(Sum('quantity'), 0),
            total_net_val=Coalesce(Sum(F('quantity') * F('net_price'), output_field=DecimalField()), Decimal('0.00'))
        )
        .order_by('-total_ordered_qty')[:10]
    )
    top_ordered_list = []
    for item in top_ordered_qs:
        p_name = item['product__name'] or item['product_name_snapshot'] or 'Produkt'
        top_ordered_list.append({
            'name': p_name,
            'quantity': item['total_ordered_qty'],
            'total_net': item['total_net_val']
        })
    top_ordered_labels = [item['name'] for item in top_ordered_list[:5]]
    top_ordered_data = [item['quantity'] for item in top_ordered_list[:5]]

    # 4. Zastosowanie filtrów dat do zużycia w zleceniach pracy (WorkOrders)
    work_orders_qs = WorkOrder.objects.filter(tenant=tenant)
    if date_from:
        work_orders_qs = work_orders_qs.filter(created_at__date__gte=date_from)
    if date_to:
        work_orders_qs = work_orders_qs.filter(created_at__date__lte=date_to)

    period_work_orders_count = work_orders_qs.count()

    consumed_qs = WorkOrderProduct.objects.filter(order__tenant=tenant)
    if date_from:
        consumed_qs = consumed_qs.filter(order__created_at__date__gte=date_from)
    if date_to:
        consumed_qs = consumed_qs.filter(order__created_at__date__lte=date_to)

    period_consumed_items_qty = consumed_qs.aggregate(total=Coalesce(Sum('quantity'), 0))['total']
    period_consumed_val_net = consumed_qs.annotate(
        val=F('quantity') * F('unit_price_net')
    ).aggregate(total=Coalesce(Sum('val'), Decimal('0.00'), output_field=DecimalField()))['total']

    top_consumed_qs = (
        consumed_qs.values('name_snapshot')
        .annotate(
            total_used_qty=Coalesce(Sum('quantity'), 0),
            total_net_val=Coalesce(Sum(F('quantity') * F('unit_price_net'), output_field=DecimalField()), Decimal('0.00'))
        )
        .order_by('-total_used_qty')[:10]
    )
    top_consumed_list = []
    for item in top_consumed_qs:
        if item['name_snapshot']:
            top_consumed_list.append({
                'name': item['name_snapshot'],
                'quantity': item['total_used_qty'],
                'total_net': item['total_net_val']
            })
    top_consumed_labels = [item['name'] for item in top_consumed_list[:5]]
    top_consumed_data = [item['quantity'] for item in top_consumed_list[:5]]

    # 5. Wykres: Wartość magazynu wg lokalizacji (stan bieżący)
    location_values = (
        batches_with_val.values('location__name')
        .annotate(loc_val=Coalesce(Sum('batch_val_net'), Decimal('0.00'), output_field=DecimalField()))
        .order_by('-loc_val')
    )
    location_labels = [item['location__name'] or 'Domyślny' for item in location_values]
    location_data = [float(item['loc_val'] or 0) for item in location_values]

    # 6. Statusy zamówień w oknie czasowym
    orders_by_status = (
        orders_qs.values('status')
        .annotate(cnt=Count('id'))
    )
    status_dict = {
        'CREATED': 'Szkic / Utworzone',
        'ORDERED': 'Wysłane do dostawcy',
        'COMPLETED': 'Zakończone / Dostarczone',
        'CANCELLED': 'Anulowane'
    }
    order_status_labels = [status_dict.get(item['status'], item['status']) for item in orders_by_status]
    order_status_data = [item['cnt'] for item in orders_by_status]

    # 7. Wykres liniowy trendu dla wybranego produktu: Odebrane zamówienia vs Zużycie
    tenant_products = Product.objects.filter(tenant=tenant).order_by('-is_favourite', 'name')
    tenant_locations = Location.objects.filter(tenant=tenant).order_by('-is_default', 'name')

    selected_product_id = request.GET.get('product_id', '').strip()
    selected_location_id = request.GET.get('location_id', '').strip()

    selected_product = None
    if selected_product_id:
        selected_product = tenant_products.filter(id=selected_product_id).first()
    if not selected_product and tenant_products.exists():
        selected_product = tenant_products.first()

    selected_location = None
    if selected_location_id:
        selected_location = tenant_locations.filter(id=selected_location_id).first()

    trend_dates = []
    trend_received_data = []
    trend_consumed_data = []
    trend_total_received = 0
    trend_total_consumed = 0
    trend_net_balance = 0

    if selected_product:
        # A. Odebrane zamówienia (status='COMPLETED')
        trend_orders_qs = Order.objects.filter(
            tenant=tenant,
            product=selected_product,
            status='COMPLETED'
        )
        if selected_location:
            trend_orders_qs = trend_orders_qs.filter(
                productbatch__location=selected_location
            ).distinct()

        if date_from:
            trend_orders_qs = trend_orders_qs.filter(updated_at__date__gte=date_from)
        if date_to:
            trend_orders_qs = trend_orders_qs.filter(updated_at__date__lte=date_to)

        orders_by_date = (
            trend_orders_qs
            .values('updated_at__date')
            .annotate(qty=Coalesce(Sum('quantity'), 0))
            .order_by('updated_at__date')
        )
        orders_date_map = {
            item['updated_at__date'].strftime('%Y-%m-%d'): item['qty']
            for item in orders_by_date if item['updated_at__date']
        }

        # B. Zużycie w zleceniach pracy (WorkOrderProduct)
        trend_consumed_qs = WorkOrderProduct.objects.filter(
            order__tenant=tenant
        ).filter(
            Q(product_batch__product=selected_product) | Q(name_snapshot=selected_product.name)
        )
        if selected_location:
            trend_consumed_qs = trend_consumed_qs.filter(
                product_batch__location=selected_location
            )

        if date_from:
            trend_consumed_qs = trend_consumed_qs.filter(order__created_at__date__gte=date_from)
        if date_to:
            trend_consumed_qs = trend_consumed_qs.filter(order__created_at__date__lte=date_to)

        consumed_by_date = (
            trend_consumed_qs
            .values('order__created_at__date')
            .annotate(qty=Coalesce(Sum('quantity'), 0))
            .order_by('order__created_at__date')
        )
        consumed_date_map = {
            item['order__created_at__date'].strftime('%Y-%m-%d'): item['qty']
            for item in consumed_by_date if item['order__created_at__date']
        }

        # C. Budowanie ciągłej osi dat
        all_active_date_strs = list(orders_date_map.keys()) + list(consumed_date_map.keys())
        today = timezone.now().date()

        if date_from and date_to:
            start_d = date_from
            end_d = date_to
        elif date_from and not date_to:
            start_d = date_from
            end_d = today
        elif date_to and not date_from:
            start_d = date_to - datetime.timedelta(days=30)
            end_d = date_to
        else:
            if all_active_date_strs:
                active_dates = [datetime.datetime.strptime(d, '%Y-%m-%d').date() for d in all_active_date_strs]
                min_d = min(active_dates)
                max_d = max(active_dates)
                start_d = min_d
                end_d = max_d
            else:
                end_d = today
                start_d = today - datetime.timedelta(days=14)

        if start_d <= end_d:
            cur_d = start_d
            while cur_d <= end_d:
                d_str = cur_d.strftime('%Y-%m-%d')
                rec_val = int(orders_date_map.get(d_str, 0))
                con_val = int(consumed_date_map.get(d_str, 0))
                trend_dates.append(d_str)
                trend_received_data.append(rec_val)
                trend_consumed_data.append(con_val)
                cur_d += datetime.timedelta(days=1)

        trend_total_received = sum(trend_received_data)
        trend_total_consumed = sum(trend_consumed_data)
        trend_net_balance = trend_total_received - trend_total_consumed

    context = {
        # Filtry dat
        'date_from': date_from.strftime('%Y-%m-%d') if date_from else '',
        'date_to': date_to.strftime('%Y-%m-%d') if date_to else '',
        'preset': preset,
        'has_date_filter': bool(date_from or date_to or preset),

        # Bieżące KPI magazynu
        'total_products': total_products,
        'total_stock_items': total_stock_items,
        'total_inventory_net': total_inventory_net,
        'total_inventory_gross': total_inventory_gross,
        'low_stock_products_count': low_stock_products_count,
        'estimated_restock_cost': estimated_restock_cost,

        # KPI okresowe (w wybranym oknie)
        'period_orders_count': period_orders_count,
        'period_ordered_items_qty': period_ordered_items_qty,
        'period_ordered_val_net': period_ordered_val_net,
        'period_work_orders_count': period_work_orders_count,
        'period_consumed_items_qty': period_consumed_items_qty,
        'period_consumed_val_net': period_consumed_val_net,

        # Wykres liniowy trendów produktu
        'tenant_products': tenant_products,
        'tenant_locations': tenant_locations,
        'selected_product': selected_product,
        'selected_location': selected_location,
        'trend_dates': trend_dates,
        'trend_received_data': trend_received_data,
        'trend_consumed_data': trend_consumed_data,
        'trend_total_received': trend_total_received,
        'trend_total_consumed': trend_total_consumed,
        'trend_net_balance': trend_net_balance,

        # Dane do Chart.js
        'top_ordered_labels': top_ordered_labels,
        'top_ordered_data': top_ordered_data,
        'top_ordered_list': top_ordered_list,
        'top_consumed_labels': top_consumed_labels,
        'top_consumed_data': top_consumed_data,
        'top_consumed_list': top_consumed_list,
        'location_labels': location_labels,
        'location_data': location_data,
        'order_status_labels': order_status_labels,
        'order_status_data': order_status_data,
    }
    return render(request, 'analytics/dashboard.html', context)


@login_required
def restock_cost_report(request):
    """Raport przewidywanych kosztów zamówień na podstawie stanów minimalnych."""
    tenant = request.user.tenant
    location_id = request.GET.get('location')
    locations = Location.objects.filter(tenant=tenant)

    selected_location = None
    if location_id:
        selected_location = locations.filter(id=location_id).first()

    report_data = _calculate_restock_report_data(tenant, location_id=location_id if selected_location else None)

    context = {
        'locations': locations,
        'selected_location': selected_location,
        'report': report_data,
    }
    return render(request, 'analytics/restock_cost_report.html', context)


@login_required
def export_stock_csv(request):
    """Eksport aktualnego stanu magazynowego do pliku CSV."""
    tenant = request.user.tenant
    location_id = request.GET.get('location')

    products = Product.objects.filter(tenant=tenant).prefetch_related('batches__location', 'supplier_mappings__supplier')
    if location_id:
        products = products.filter(batches__location_id=location_id).distinct()

    response = HttpResponse(content_type='text/csv; charset=utf-8-sig')
    timestamp = timezone.now().strftime("%Y-%m-%d_%H%M")
    response['Content-Disposition'] = f'attachment; filename="stan_magazynowy_{timestamp}.csv"'

    writer = csv.writer(response, delimiter=';')
    writer.writerow([
        'ID Produktu',
        'Numer lokalny',
        'Nazwa produktu',
        'Magazyn',
        'Ilość w partii',
        'Cena zakupu netto (zł)',
        'Cena zakupu brutto (zł)',
        'Wartość partii netto (zł)',
        'Wartość partii brutto (zł)',
        'Próg alarmowy',
        'Stawka VAT (%)',
        'Główny dostawca'
    ])

    for prod in products:
        batches = prod.batches.filter(tenant=tenant)
        if location_id:
            batches = batches.filter(location_id=location_id)

        primary_supplier = prod.supplier_mappings.filter(tenant=tenant).first()
        supplier_name = primary_supplier.supplier.name if primary_supplier else ""

        if not batches.exists():
            writer.writerow([
                str(prod.id),
                prod.local_id or '',
                prod.name,
                '-',
                0,
                '0.00',
                '0.00',
                '0.00',
                '0.00',
                prod.min_threshold,
                prod.vat_rate,
                supplier_name
            ])
        else:
            for b in batches:
                val_net = (Decimal(b.current_stock or 0) * (b.net_price or Decimal('0.00'))).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
                val_gross = (Decimal(b.current_stock or 0) * (b.gross_price or Decimal('0.00'))).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
                loc_name = b.location.name if b.location else 'Domyślny'
                writer.writerow([
                    str(prod.id),
                    prod.local_id or '',
                    prod.name,
                    loc_name,
                    b.current_stock or 0,
                    f"{b.net_price:.2f}",
                    f"{b.gross_price:.2f}",
                    f"{val_net:.2f}",
                    f"{val_gross:.2f}",
                    prod.min_threshold,
                    prod.vat_rate,
                    supplier_name
                ])

    return response


@login_required
def export_stock_excel(request):
    """Eksport aktualnego stanu magazynowego do pliku Excel (.xlsx)."""
    tenant = request.user.tenant
    location_id = request.GET.get('location')

    products = Product.objects.filter(tenant=tenant).prefetch_related('batches__location', 'supplier_mappings__supplier')
    if location_id:
        products = products.filter(batches__location_id=location_id).distinct()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Stan magazynowy"

    # Style
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=16, bold=True, color="1F4E78")
    meta_font = Font(name="Calibri", size=10, italic=True, color="595959")
    bold_font = Font(name="Calibri", size=11, bold=True)
    border_thin = Side(border_style="thin", color="D9D9D9")
    cell_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)

    ws['A1'] = f"Raport stanu magazynowego - {tenant.name}"
    ws['A1'].font = title_font
    ws['A2'] = f"Wygenerowano: {timezone.now().strftime('%Y-%m-%d %H:%M:%S')}"
    ws['A2'].font = meta_font

    headers = [
        "Lp.", "Produkt", "Nr", "Magazyn", "Ilość",
        "Cena netto (zł)", "Cena brutto (zł)", "Wartość netto (zł)", "Wartość brutto (zł)",
        "Próg min.", "VAT %", "Główny dostawca"
    ]

    row_num = 4
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=row_num, column=col_num, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    row_num = 5
    counter = 1
    total_qty = 0
    total_val_net = Decimal('0.00')
    total_val_gross = Decimal('0.00')

    for prod in products:
        batches = prod.batches.filter(tenant=tenant)
        if location_id:
            batches = batches.filter(location_id=location_id)

        primary_supplier = prod.supplier_mappings.filter(tenant=tenant).first()
        supplier_name = primary_supplier.supplier.name if primary_supplier else ""

        if not batches.exists():
            data_row = [
                counter, prod.name, prod.local_id or "", "-", 0,
                0.0, 0.0, 0.0, 0.0,
                prod.min_threshold, f"{prod.vat_rate}%", supplier_name
            ]
            for col_idx, val in enumerate(data_row, 1):
                c = ws.cell(row=row_num, column=col_idx, value=val)
                c.border = cell_border
                if col_idx in [5, 10]:
                    c.alignment = Alignment(horizontal="right")
                elif col_idx in [6, 7, 8, 9]:
                    c.number_format = '#,##0.00'
                    c.alignment = Alignment(horizontal="right")
            row_num += 1
            counter += 1
        else:
            for b in batches:
                v_net = (Decimal(b.current_stock or 0) * (b.net_price or Decimal('0.00'))).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
                v_gross = (Decimal(b.current_stock or 0) * (b.gross_price or Decimal('0.00'))).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
                loc_name = b.location.name if b.location else 'Domyślny'

                total_qty += b.current_stock or 0
                total_val_net += v_net
                total_val_gross += v_gross

                data_row = [
                    counter, prod.name, prod.local_id or "", loc_name, b.current_stock or 0,
                    float(b.net_price or 0), float(b.gross_price or 0), float(v_net), float(v_gross),
                    prod.min_threshold, f"{prod.vat_rate}%", supplier_name
                ]
                for col_idx, val in enumerate(data_row, 1):
                    c = ws.cell(row=row_num, column=col_idx, value=val)
                    c.border = cell_border
                    if col_idx in [5, 10]:
                        c.alignment = Alignment(horizontal="right")
                    elif col_idx in [6, 7, 8, 9]:
                        c.number_format = '#,##0.00'
                        c.alignment = Alignment(horizontal="right")
                row_num += 1
                counter += 1

    summary_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    ws.cell(row=row_num, column=2, value="ŁĄCZNIE").font = bold_font
    c_qty = ws.cell(row=row_num, column=5, value=total_qty)
    c_qty.font = bold_font
    c_qty.alignment = Alignment(horizontal="right")

    c_net = ws.cell(row=row_num, column=8, value=float(total_val_net))
    c_net.font = bold_font
    c_net.number_format = '#,##0.00'
    c_gross = ws.cell(row=row_num, column=9, value=float(total_val_gross))
    c_gross.font = bold_font
    c_gross.number_format = '#,##0.00'

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=row_num, column=col_idx)
        cell.fill = summary_fill
        cell.border = cell_border

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row < 4:
                continue
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    timestamp = timezone.now().strftime("%Y-%m-%d_%H%M")
    response['Content-Disposition'] = f'attachment; filename="stan_magazynowy_{timestamp}.xlsx"'
    wb.save(response)
    return response


@login_required
def export_restock_csv(request):
    """Eksport raportu prognozowanych kosztów do CSV."""
    tenant = request.user.tenant
    location_id = request.GET.get('location')
    report_data = _calculate_restock_report_data(tenant, location_id=location_id)

    response = HttpResponse(content_type='text/csv; charset=utf-8-sig')
    timestamp = timezone.now().strftime("%Y-%m-%d_%H%M")
    response['Content-Disposition'] = f'attachment; filename="prognoza_kosztow_zamowien_{timestamp}.csv"'

    writer = csv.writer(response, delimiter=';')
    writer.writerow([
        'Produkt',
        'Stan aktualny',
        'Stan minimalny',
        'Brakująca ilość',
        'Szacowana cena netto (zł)',
        'Szacowana cena brutto (zł)',
        'Szacowany koszt netto (zł)',
        'Szacowany koszt brutto (zł)',
        'Dostawca'
    ])

    for item in report_data['items']:
        writer.writerow([
            item['product'].name,
            item['current_stock'],
            item['min_stock_level'],
            item['shortage_quantity'],
            f"{item['unit_price_net']:.2f}",
            f"{item['unit_price_gross']:.2f}",
            f"{item['total_cost_net']:.2f}",
            f"{item['total_cost_gross']:.2f}",
            item['supplier_name']
        ])

    return response


@login_required
def export_restock_excel(request):
    """Eksport raportu prognozowanych kosztów do Excel (.xlsx)."""
    tenant = request.user.tenant
    location_id = request.GET.get('location')
    report_data = _calculate_restock_report_data(tenant, location_id=location_id)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Prognoza kosztów zamówień"

    header_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=16, bold=True, color="2F5597")
    meta_font = Font(name="Calibri", size=10, italic=True, color="595959")
    bold_font = Font(name="Calibri", size=11, bold=True)
    border_thin = Side(border_style="thin", color="D9D9D9")
    cell_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)

    ws['A1'] = f"Raport prognozy kosztów zamówień (uzupełnienie braków) - {tenant.name}"
    ws['A1'].font = title_font
    ws['A2'] = f"Wygenerowano: {timezone.now().strftime('%Y-%m-%d %H:%M:%S')}"
    ws['A2'].font = meta_font

    headers = [
        "Lp.", "Produkt", "Stan aktualny", "Stan min.", "Braki (szt.)",
        "Cena jedn. netto (zł)", "Cena jedn. brutto (zł)",
        "Koszt netto (zł)", "Koszt brutto (zł)", "Sugerowany dostawca"
    ]

    row_num = 4
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=row_num, column=col_num, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    row_num = 5
    for idx, item in enumerate(report_data['items'], 1):
        data_row = [
            idx, item['product'].name, item['current_stock'], item['min_stock_level'],
            item['shortage_quantity'], float(item['unit_price_net']), float(item['unit_price_gross']),
            float(item['total_cost_net']), float(item['total_cost_gross']), item['supplier_name']
        ]
        for col_idx, val in enumerate(data_row, 1):
            c = ws.cell(row=row_num, column=col_idx, value=val)
            c.border = cell_border
            if col_idx in [3, 4, 5]:
                c.alignment = Alignment(horizontal="right")
            elif col_idx in [6, 7, 8, 9]:
                c.number_format = '#,##0.00'
                c.alignment = Alignment(horizontal="right")
        row_num += 1

    summary_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    ws.cell(row=row_num, column=2, value="ŁĄCZNIE").font = bold_font
    c_short = ws.cell(row=row_num, column=5, value=report_data['total_shortage_qty'])
    c_short.font = bold_font
    c_short.alignment = Alignment(horizontal="right")

    c_net = ws.cell(row=row_num, column=8, value=float(report_data['total_estimated_cost_net']))
    c_net.font = bold_font
    c_net.number_format = '#,##0.00'

    c_gross = ws.cell(row=row_num, column=9, value=float(report_data['total_estimated_cost_gross']))
    c_gross.font = bold_font
    c_gross.number_format = '#,##0.00'

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=row_num, column=col_idx)
        cell.fill = summary_fill
        cell.border = cell_border

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row < 4:
                continue
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    timestamp = timezone.now().strftime("%Y-%m-%d_%H%M")
    response['Content-Disposition'] = f'attachment; filename="prognoza_kosztow_zamowien_{timestamp}.xlsx"'
    wb.save(response)
    return response


@login_required
def export_consumption_csv(request):
    """Eksport raportu zużycia materiałów w zleceniach pracy do CSV."""
    tenant = request.user.tenant
    date_from, date_to, _ = _parse_date_range(request)

    consumed_qs = WorkOrderProduct.objects.filter(order__tenant=tenant).select_related('order')
    if date_from:
        consumed_qs = consumed_qs.filter(order__created_at__date__gte=date_from)
    if date_to:
        consumed_qs = consumed_qs.filter(order__created_at__date__lte=date_to)

    consumed_qs = consumed_qs.order_by('-order__created_at')

    response = HttpResponse(content_type='text/csv; charset=utf-8-sig')
    timestamp = timezone.now().strftime("%Y-%m-%d_%H%M")
    response['Content-Disposition'] = f'attachment; filename="zuzycie_materialow_{timestamp}.csv"'

    writer = csv.writer(response, delimiter=';')
    writer.writerow([
        'Data zlecenia',
        'Zlecenie',
        'Klient',
        'Produkt / Część',
        'Ilość',
        'Cena jedn. netto (zł)',
        'Wartość netto (zł)',
        'Stawka VAT (%)',
        'Wartość brutto (zł)'
    ])

    for item in consumed_qs:
        val_net = (Decimal(item.quantity) * item.unit_price_net).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
        vat_factor = 1 + (item.vat_rate / Decimal('100'))
        val_gross = (val_net * vat_factor).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
        writer.writerow([
            item.order.created_at.strftime('%Y-%m-%d %H:%M'),
            item.order.name,
            item.order.client,
            item.name_snapshot,
            item.quantity,
            f"{item.unit_price_net:.2f}",
            f"{val_net:.2f}",
            f"{item.vat_rate:.2f}",
            f"{val_gross:.2f}"
        ])

    return response


@login_required
def export_consumption_excel(request):
    """Eksport raportu zużycia materiałów w zleceniach pracy do Excel (.xlsx)."""
    tenant = request.user.tenant
    date_from, date_to, _ = _parse_date_range(request)

    consumed_qs = WorkOrderProduct.objects.filter(order__tenant=tenant).select_related('order')
    if date_from:
        consumed_qs = consumed_qs.filter(order__created_at__date__gte=date_from)
    if date_to:
        consumed_qs = consumed_qs.filter(order__created_at__date__lte=date_to)

    consumed_qs = consumed_qs.order_by('-order__created_at')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Zużycie materiałów"

    header_fill = PatternFill(start_color="D35400", end_color="D35400", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=16, bold=True, color="D35400")
    meta_font = Font(name="Calibri", size=10, italic=True, color="595959")
    bold_font = Font(name="Calibri", size=11, bold=True)
    border_thin = Side(border_style="thin", color="D9D9D9")
    cell_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)

    period_info = ""
    if date_from and date_to:
        period_info = f" za okres {date_from} do {date_to}"
    elif date_from:
        period_info = f" od {date_from}"
    elif date_to:
        period_info = f" do {date_to}"

    ws['A1'] = f"Raport zużycia materiałów w zleceniach{period_info} - {tenant.name}"
    ws['A1'].font = title_font
    ws['A2'] = f"Wygenerowano: {timezone.now().strftime('%Y-%m-%d %H:%M:%S')}"
    ws['A2'].font = meta_font

    headers = [
        "Lp.", "Data zlecenia", "Zlecenie", "Klient", "Produkt / Część", "Ilość",
        "Cena jedn. netto (zł)", "Wartość netto (zł)", "VAT (%)", "Wartość brutto (zł)"
    ]

    row_num = 4
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=row_num, column=col_num, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    row_num = 5
    total_qty = 0
    total_val_net = Decimal('0.00')
    total_val_gross = Decimal('0.00')

    for idx, item in enumerate(consumed_qs, 1):
        v_net = (Decimal(item.quantity) * item.unit_price_net).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
        vat_factor = 1 + (item.vat_rate / Decimal('100'))
        v_gross = (v_net * vat_factor).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

        total_qty += item.quantity
        total_val_net += v_net
        total_val_gross += v_gross

        data_row = [
            idx, item.order.created_at.strftime('%Y-%m-%d %H:%M'), item.order.name, item.order.client,
            item.name_snapshot, item.quantity, float(item.unit_price_net), float(v_net),
            float(item.vat_rate), float(v_gross)
        ]
        for col_idx, val in enumerate(data_row, 1):
            c = ws.cell(row=row_num, column=col_idx, value=val)
            c.border = cell_border
            if col_idx in [6]:
                c.alignment = Alignment(horizontal="right")
            elif col_idx in [7, 8, 9, 10]:
                c.number_format = '#,##0.00'
                c.alignment = Alignment(horizontal="right")
        row_num += 1

    summary_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    ws.cell(row=row_num, column=3, value="ŁĄCZNIE").font = bold_font
    c_qty = ws.cell(row=row_num, column=6, value=total_qty)
    c_qty.font = bold_font
    c_qty.alignment = Alignment(horizontal="right")

    c_net = ws.cell(row=row_num, column=8, value=float(total_val_net))
    c_net.font = bold_font
    c_net.number_format = '#,##0.00'

    c_gross = ws.cell(row=row_num, column=10, value=float(total_val_gross))
    c_gross.font = bold_font
    c_gross.number_format = '#,##0.00'

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=row_num, column=col_idx)
        cell.fill = summary_fill
        cell.border = cell_border

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row < 4:
                continue
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    timestamp = timezone.now().strftime("%Y-%m-%d_%H%M")
    response['Content-Disposition'] = f'attachment; filename="zuzycie_materialow_{timestamp}.xlsx"'
    wb.save(response)
    return response


@login_required
def export_orders_csv(request):
    """Eksport raportu zamówień w wybranym oknie czasowym do CSV."""
    tenant = request.user.tenant
    date_from, date_to, _ = _parse_date_range(request)

    orders_qs = Order.objects.filter(tenant=tenant).select_related('supplier', 'product')
    if date_from:
        orders_qs = orders_qs.filter(created_at__date__gte=date_from)
    if date_to:
        orders_qs = orders_qs.filter(created_at__date__lte=date_to)

    orders_qs = orders_qs.order_by('-created_at')

    response = HttpResponse(content_type='text/csv; charset=utf-8-sig')
    timestamp = timezone.now().strftime("%Y-%m-%d_%H%M")
    response['Content-Disposition'] = f'attachment; filename="zamowienia_{timestamp}.csv"'

    writer = csv.writer(response, delimiter=';')
    writer.writerow([
        'Data utworzenia',
        'Nr zamówienia',
        'Produkt',
        'Dostawca',
        'Ilość',
        'Cena netto (zł)',
        'Wartość netto (zł)',
        'Cena brutto (zł)',
        'Wartość brutto (zł)',
        'Status',
        'Typ'
    ])

    for o in orders_qs:
        p_name = o.product.name if o.product else o.product_name_snapshot
        s_name = o.supplier.name if o.supplier else ''
        val_net = (Decimal(o.quantity) * (o.net_price or Decimal('0.00'))).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
        val_gross = (Decimal(o.quantity) * (o.gross_price or Decimal('0.00'))).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

        writer.writerow([
            o.created_at.strftime('%Y-%m-%d %H:%M'),
            o.order_number or '',
            p_name,
            s_name,
            o.quantity,
            f"{o.net_price:.2f}" if o.net_price else "0.00",
            f"{val_net:.2f}",
            f"{o.gross_price:.2f}" if o.gross_price else "0.00",
            f"{val_gross:.2f}",
            o.get_status_display(),
            o.get_order_type_display()
        ])

    return response


@login_required
def export_orders_excel(request):
    """Eksport raportu zamówień w wybranym oknie czasowym do Excel (.xlsx)."""
    tenant = request.user.tenant
    date_from, date_to, _ = _parse_date_range(request)

    orders_qs = Order.objects.filter(tenant=tenant).select_related('supplier', 'product')
    if date_from:
        orders_qs = orders_qs.filter(created_at__date__gte=date_from)
    if date_to:
        orders_qs = orders_qs.filter(created_at__date__lte=date_to)

    orders_qs = orders_qs.order_by('-created_at')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Zamówienia"

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=16, bold=True, color="1F4E78")
    meta_font = Font(name="Calibri", size=10, italic=True, color="595959")
    bold_font = Font(name="Calibri", size=11, bold=True)
    border_thin = Side(border_style="thin", color="D9D9D9")
    cell_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)

    period_info = ""
    if date_from and date_to:
        period_info = f" za okres {date_from} do {date_to}"
    elif date_from:
        period_info = f" od {date_from}"
    elif date_to:
        period_info = f" do {date_to}"

    ws['A1'] = f"Raport zamówień towaru{period_info} - {tenant.name}"
    ws['A1'].font = title_font
    ws['A2'] = f"Wygenerowano: {timezone.now().strftime('%Y-%m-%d %H:%M:%S')}"
    ws['A2'].font = meta_font

    headers = [
        "Lp.", "Data utworzenia", "Nr", "Produkt", "Dostawca", "Ilość",
        "Cena netto (zł)", "Wartość netto (zł)", "Cena brutto (zł)", "Wartość brutto (zł)", "Status", "Typ"
    ]

    row_num = 4
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=row_num, column=col_num, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    row_num = 5
    total_qty = 0
    total_val_net = Decimal('0.00')
    total_val_gross = Decimal('0.00')

    for idx, o in enumerate(orders_qs, 1):
        p_name = o.product.name if o.product else o.product_name_snapshot
        s_name = o.supplier.name if o.supplier else ''
        v_net = (Decimal(o.quantity) * (o.net_price or Decimal('0.00'))).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
        v_gross = (Decimal(o.quantity) * (o.gross_price or Decimal('0.00'))).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

        total_qty += o.quantity
        total_val_net += v_net
        total_val_gross += v_gross

        data_row = [
            idx, o.created_at.strftime('%Y-%m-%d %H:%M'), o.order_number or "", p_name, s_name,
            o.quantity, float(o.net_price or 0), float(v_net), float(o.gross_price or 0), float(v_gross),
            o.get_status_display(), o.get_order_type_display()
        ]
        for col_idx, val in enumerate(data_row, 1):
            c = ws.cell(row=row_num, column=col_idx, value=val)
            c.border = cell_border
            if col_idx in [6]:
                c.alignment = Alignment(horizontal="right")
            elif col_idx in [7, 8, 9, 10]:
                c.number_format = '#,##0.00'
                c.alignment = Alignment(horizontal="right")
        row_num += 1

    summary_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    ws.cell(row=row_num, column=4, value="ŁĄCZNIE").font = bold_font
    c_qty = ws.cell(row=row_num, column=6, value=total_qty)
    c_qty.font = bold_font
    c_qty.alignment = Alignment(horizontal="right")

    c_net = ws.cell(row=row_num, column=8, value=float(total_val_net))
    c_net.font = bold_font
    c_net.number_format = '#,##0.00'

    c_gross = ws.cell(row=row_num, column=10, value=float(total_val_gross))
    c_gross.font = bold_font
    c_gross.number_format = '#,##0.00'

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=row_num, column=col_idx)
        cell.fill = summary_fill
        cell.border = cell_border

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row < 4:
                continue
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    timestamp = timezone.now().strftime("%Y-%m-%d_%H%M")
    response['Content-Disposition'] = f'attachment; filename="zamowienia_{timestamp}.xlsx"'
    wb.save(response)
    return response
