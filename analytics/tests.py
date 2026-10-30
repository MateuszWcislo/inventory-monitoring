from decimal import Decimal
import datetime
import io
import openpyxl

from django.test import TestCase, Client
from django.urls import reverse
from django.utils import timezone

from tenants.models import Tenant
from users.models import User
from inventory.models import Product, ProductBatch, Location
from suppliers.models import Supplier
from orders.models import Order
from work_orders.models import WorkOrder, WorkOrderProduct


class AnalyticsAndReportingTests(TestCase):
    def setUp(self):
        # Tenant 1
        self.tenant1 = Tenant.objects.create(name="Firma Testowa 1", subdomain="firma1")
        self.user1 = User.objects.create_user(
            username="user1",
            email="user1@firma1.pl",
            password="password123",
            tenant=self.tenant1,
            role="ADMIN"
        )

        # Tenant 2
        self.tenant2 = Tenant.objects.create(name="Firma Testowa 2", subdomain="firma2")
        self.user2 = User.objects.create_user(
            username="user2",
            email="user2@firma2.pl",
            password="password123",
            tenant=self.tenant2,
            role="ADMIN"
        )

        self.client1 = Client()
        self.client1.login(username="user1", password="password123")

        self.client2 = Client()
        self.client2.login(username="user2", password="password123")

        # Dane dla Tenanta 1
        self.loc1 = Location.objects.create(tenant=self.tenant1, name="Magazyn Główny", is_default=True)
        self.loc2 = Location.objects.create(tenant=self.tenant1, name="Magazyn Pomocniczy", is_default=False)

        self.supplier1 = Supplier.objects.create(tenant=self.tenant1, name="Dostawca Alfa")

        # Produkt 1: niski stan (brak: 10 - 2 = 8 szt.), cena netto 50.00 zł
        self.p1 = Product.objects.create(
            tenant=self.tenant1,
            name="Olej silnikowy 5W30",
            min_threshold=10,
            vat_rate=23
        )
        self.b1 = ProductBatch(
            tenant=self.tenant1,
            product=self.p1,
            location=self.loc1,
            current_stock=2,
            net_price=Decimal('50.00'),
            gross_price=Decimal('61.50')
        )
        self.b1._skip_signal = True
        self.b1.save()

        # Produkt 2: stan prawidłowy (15 >= 5)
        self.p2 = Product.objects.create(
            tenant=self.tenant1,
            name="Filtr powietrza",
            min_threshold=5,
            vat_rate=23
        )
        self.b2 = ProductBatch(
            tenant=self.tenant1,
            product=self.p2,
            location=self.loc2,
            current_stock=15,
            net_price=Decimal('20.00'),
            gross_price=Decimal('24.60')
        )
        self.b2._skip_signal = True
        self.b2.save()

        # Czyścimy ewentualne auto-zamówienia
        Order.objects.filter(tenant=self.tenant1).delete()

        # Zamówienia z różnymi datami
        now = timezone.now()
        self.order_recent = Order.objects.create(
            tenant=self.tenant1,
            product=self.p1,
            supplier=self.supplier1,
            quantity=10,
            net_price=Decimal('50.00'),
            gross_price=Decimal('61.50'),
            status='ORDERED'
        )

        self.order_old = Order(
            tenant=self.tenant1,
            product=self.p2,
            supplier=self.supplier1,
            quantity=5,
            net_price=Decimal('20.00'),
            gross_price=Decimal('24.60'),
            status='ORDERED'
        )
        self.order_old._skip_signal = True
        self.order_old.save()
        # Ustawienie starszej daty dla order_old
        Order.objects.filter(id=self.order_old.id).update(created_at=now - datetime.timedelta(days=60))

        # Zlecenia pracy (Work Orders) i zużycie
        self.wo_recent = WorkOrder.objects.create(
            tenant=self.tenant1,
            name="Wymiana oleju i filtrów",
            client="Klient A",
            status="COMPLETED"
        )
        self.wop_recent = WorkOrderProduct.objects.create(
            order=self.wo_recent,
            product_batch=self.b1,
            name_snapshot="Olej silnikowy 5W30",
            unit_price_net=Decimal('50.00'),
            vat_rate=Decimal('23.00'),
            quantity=4
        )

        self.wo_old = WorkOrder.objects.create(
            tenant=self.tenant1,
            name="Stare zlecenie",
            client="Klient B",
            status="COMPLETED"
        )
        self.wop_old = WorkOrderProduct.objects.create(
            order=self.wo_old,
            product_batch=self.b2,
            name_snapshot="Filtr powietrza",
            unit_price_net=Decimal('20.00'),
            vat_rate=Decimal('23.00'),
            quantity=2
        )
        WorkOrder.objects.filter(id=self.wo_old.id).update(created_at=now - datetime.timedelta(days=60))

        # Dane dla Tenanta 2 (izolacja)
        self.p_tenant2 = Product.objects.create(
            tenant=self.tenant2,
            name="Produkt Tenanta 2",
            min_threshold=100,
            vat_rate=23
        )
        self.b_tenant2 = ProductBatch(
            tenant=self.tenant2,
            product=self.p_tenant2,
            current_stock=1,
            net_price=Decimal('100.00'),
            gross_price=Decimal('123.00')
        )
        self.b_tenant2._skip_signal = True
        self.b_tenant2.save()

    def test_analytics_dashboard_metrics_and_isolation(self):
        """Weryfikacja podsumowań KPI i izolacji danych tenanta w dashboardzie."""
        response = self.client1.get(reverse('analytics_dashboard'))
        self.assertEqual(response.status_code, 200)

        # Tenant 1 ma 2 produkty
        self.assertEqual(response.context['total_products'], 2)
        # Suma sztuk: 2 + 15 = 17
        self.assertEqual(response.context['total_stock_items'], 17)
        # Wartość netto: 2*50 + 15*20 = 100 + 300 = 400.00 zł
        self.assertEqual(Decimal(str(response.context['total_inventory_net'])), Decimal('400.00'))
        # Pozycje poniżej min: tylko p1
        self.assertEqual(response.context['low_stock_products_count'], 1)

    def test_analytics_dashboard_date_filtering(self):
        """Weryfikacja filtrów czasowych (data od - data do oraz presety) w dashboardzie."""
        today = timezone.now().date()
        date_from_recent = (today - datetime.timedelta(days=7)).strftime('%Y-%m-%d')
        date_to_today = today.strftime('%Y-%m-%d')

        # 1. Filtrowanie na ostatnie 7 dni (powinno uwzględnić tylko order_recent i wo_recent)
        response = self.client1.get(reverse('analytics_dashboard') + f'?date_from={date_from_recent}&date_to={date_to_today}')
        self.assertEqual(response.status_code, 200)

        self.assertEqual(response.context['period_orders_count'], 1)
        self.assertEqual(response.context['period_ordered_items_qty'], 10)
        self.assertEqual(response.context['period_work_orders_count'], 1)
        self.assertEqual(response.context['period_consumed_items_qty'], 4)
        self.assertEqual(Decimal(str(response.context['period_consumed_val_net'])), Decimal('200.00'))

        # 2. Preset 30d
        res_preset_30 = self.client1.get(reverse('analytics_dashboard') + '?preset=30d')
        self.assertEqual(res_preset_30.status_code, 200)
        self.assertEqual(res_preset_30.context['period_orders_count'], 1)

        # 3. Preset 90d (powinno uwzględnić oba zlecenia i zamówienia)
        res_preset_90 = self.client1.get(reverse('analytics_dashboard') + '?preset=90d')
        self.assertEqual(res_preset_90.status_code, 200)
        self.assertEqual(res_preset_90.context['period_orders_count'], 2)
        self.assertEqual(res_preset_90.context['period_ordered_items_qty'], 15)
        self.assertEqual(res_preset_90.context['period_work_orders_count'], 2)
        self.assertEqual(res_preset_90.context['period_consumed_items_qty'], 6)

    def test_restock_cost_report_calculation(self):
        """Weryfikacja wyliczeń braków i szacowanych kosztów zamówień."""
        response = self.client1.get(reverse('restock_cost_report'))
        self.assertEqual(response.status_code, 200)

        report = response.context['report']
        # 1 produkt z brakiem (p1: brakuje 8 szt.)
        self.assertEqual(report['total_products_count'], 1)
        self.assertEqual(report['total_shortage_qty'], 8)

        # Szacowany koszt: 8 * 50.00 = 400.00 zł netto, 8 * 61.50 = 492.00 zł brutto
        self.assertEqual(report['total_estimated_cost_net'], Decimal('400.00'))
        self.assertEqual(report['total_estimated_cost_gross'], Decimal('492.00'))

        item = report['items'][0]
        self.assertEqual(item['product'].name, "Olej silnikowy 5W30")
        self.assertEqual(item['shortage_quantity'], 8)

    def test_export_stock_csv_and_excel(self):
        """Weryfikacja generowania plików CSV i Excel stanu magazynowego."""
        # CSV
        res_csv = self.client1.get(reverse('export_stock_csv'))
        self.assertEqual(res_csv.status_code, 200)
        self.assertIn('text/csv', res_csv['Content-Type'])
        content = res_csv.content.decode('utf-8-sig')
        self.assertIn("Olej silnikowy 5W30", content)
        self.assertNotIn("Produkt Tenanta 2", content)

        # Excel
        res_xlsx = self.client1.get(reverse('export_stock_excel'))
        self.assertEqual(res_xlsx.status_code, 200)
        self.assertIn('spreadsheetml', res_xlsx['Content-Type'])
        wb = openpyxl.load_workbook(io.BytesIO(res_xlsx.content))
        ws = wb.active
        self.assertEqual(ws.title, "Stan magazynowy")
        cell_values = [cell.value for row in ws.iter_rows() for cell in row if cell.value]
        self.assertTrue(any("Olej silnikowy 5W30" in str(v) for v in cell_values))
        self.assertFalse(any("Produkt Tenanta 2" in str(v) for v in cell_values))

    def test_export_restock_csv_and_excel(self):
        """Weryfikacja generowania plików CSV i Excel raportu prognoz kosztów."""
        # CSV
        res_csv = self.client1.get(reverse('export_restock_csv'))
        self.assertEqual(res_csv.status_code, 200)
        content = res_csv.content.decode('utf-8-sig')
        self.assertIn("Olej silnikowy 5W30", content)
        self.assertIn("400.00", content)

        # Excel
        res_xlsx = self.client1.get(reverse('export_restock_excel'))
        self.assertEqual(res_xlsx.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(res_xlsx.content))
        ws = wb.active
        self.assertEqual(ws.title, "Prognoza kosztów zamówień")
        cell_values = [cell.value for row in ws.iter_rows() for cell in row if cell.value]
        self.assertTrue(any("Olej silnikowy 5W30" in str(v) for v in cell_values))
        self.assertFalse(any("Produkt Tenanta 2" in str(v) for v in cell_values))

    def test_export_consumption_and_orders_csv_and_excel(self):
        """Weryfikacja generowania plików CSV i Excel dla raportu zużycia oraz zamówień z filtrem daty."""
        today = timezone.now().date()
        date_from_recent = (today - datetime.timedelta(days=7)).strftime('%Y-%m-%d')
        date_to_today = today.strftime('%Y-%m-%d')

        # Eksport zużycia CSV z filtrem
        res_c_csv = self.client1.get(reverse('export_consumption_csv') + f'?date_from={date_from_recent}&date_to={date_to_today}')
        self.assertEqual(res_c_csv.status_code, 200)
        c_content = res_c_csv.content.decode('utf-8-sig')
        self.assertIn("Olej silnikowy 5W30", c_content)
        self.assertNotIn("Filtr powietrza", c_content)  # filtr powietrza był 60 dni temu

        # Eksport zużycia Excel z filtrem
        res_c_xlsx = self.client1.get(reverse('export_consumption_excel') + f'?date_from={date_from_recent}&date_to={date_to_today}')
        self.assertEqual(res_c_xlsx.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(res_c_xlsx.content))
        self.assertEqual(wb.active.title, "Zużycie materiałów")

        # Eksport zamówień CSV z filtrem
        res_o_csv = self.client1.get(reverse('export_orders_csv') + f'?date_from={date_from_recent}&date_to={date_to_today}')
        self.assertEqual(res_o_csv.status_code, 200)
        o_content = res_o_csv.content.decode('utf-8-sig')
        self.assertIn("Olej silnikowy 5W30", o_content)
        self.assertNotIn("Filtr powietrza", o_content)

        # Eksport zamówień Excel z filtrem
        res_o_xlsx = self.client1.get(reverse('export_orders_excel') + f'?date_from={date_from_recent}&date_to={date_to_today}')
        self.assertEqual(res_o_xlsx.status_code, 200)
        wb_o = openpyxl.load_workbook(io.BytesIO(res_o_xlsx.content))
        self.assertEqual(wb_o.active.title, "Zamówienia")

    def test_product_trend_line_chart_data(self):
        """Weryfikacja agregacji danych do wykresu liniowego trendów produktu (Odebrane vs Zużyte)."""
        today = timezone.now().date()
        date_from_str = (today - datetime.timedelta(days=7)).strftime('%Y-%m-%d')
        date_to_str = today.strftime('%Y-%m-%d')

        # Utwórz odebrane zamówienie dla p1
        order_completed = Order.objects.create(
            tenant=self.tenant1,
            product=self.p1,
            supplier=self.supplier1,
            quantity=8,
            net_price=Decimal('50.00'),
            gross_price=Decimal('61.50'),
            status='COMPLETED'
        )
        self.b1.processed_orders.add(order_completed)

        response = self.client1.get(
            reverse('analytics_dashboard') + f'?product_id={self.p1.id}&date_from={date_from_str}&date_to={date_to_str}'
        )
        self.assertEqual(response.status_code, 200)

        # Sprawdzenie kontekstu trendu
        self.assertEqual(response.context['selected_product'], self.p1)
        self.assertIn(today.strftime('%Y-%m-%d'), response.context['trend_dates'])
        self.assertEqual(response.context['trend_total_received'], 8)
        self.assertEqual(response.context['trend_total_consumed'], 4)  # z wo_recent w setUp
        self.assertEqual(response.context['trend_net_balance'], 4)  # 8 - 4 = 4

    def test_product_trend_filtering_by_location(self):
        """Weryfikacja filtrowania wykresu trendu produktu wg magazynu."""
        today = timezone.now().date()
        date_from_str = (today - datetime.timedelta(days=7)).strftime('%Y-%m-%d')
        date_to_str = today.strftime('%Y-%m-%d')

        order_completed = Order.objects.create(
            tenant=self.tenant1,
            product=self.p1,
            supplier=self.supplier1,
            quantity=12,
            net_price=Decimal('50.00'),
            gross_price=Decimal('61.50'),
            status='COMPLETED'
        )
        self.b1.processed_orders.add(order_completed)

        # Filtrowanie po loc1 (gdzie znajduje się partia b1)
        res_loc1 = self.client1.get(
            reverse('analytics_dashboard') + f'?product_id={self.p1.id}&location_id={self.loc1.id}&date_from={date_from_str}&date_to={date_to_str}'
        )
        self.assertEqual(res_loc1.status_code, 200)
        self.assertEqual(res_loc1.context['trend_total_received'], 12)
        self.assertEqual(res_loc1.context['trend_total_consumed'], 4)

        # Filtrowanie po loc2 (inny magazyn - brak partii dla p1)
        res_loc2 = self.client1.get(
            reverse('analytics_dashboard') + f'?product_id={self.p1.id}&location_id={self.loc2.id}&date_from={date_from_str}&date_to={date_to_str}'
        )
        self.assertEqual(res_loc2.status_code, 200)
        self.assertEqual(res_loc2.context['trend_total_received'], 0)
        self.assertEqual(res_loc2.context['trend_total_consumed'], 0)

    def test_product_trend_cross_tenant_isolation(self):
        """Weryfikacja izolacji danych trendów między tenantami."""
        # Użytkownik tenanta 2 próbuje podejrzeć produkt tenanta 1
        res = self.client2.get(reverse('analytics_dashboard') + f'?product_id={self.p1.id}')
        self.assertEqual(res.status_code, 200)
        # Nie powinien otrzymać produktu tenanta 1 jako wybranego
        self.assertNotEqual(res.context.get('selected_product'), self.p1)
